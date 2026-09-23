"""
parse_excel.py  —  从 marketer 填写的 Excel 解析出原始快照，生成 dashboard_data.json。
这是"手动 Excel"数据路径（在还没接通 Meta API 前使用）。
用法:  python parse_excel.py "All Meta Ads Reports.xlsx"
"""
import openpyxl, json, re, sys, os
from process import build_dashboard_data

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "All Meta Ads Reports.xlsx")

def norm(s): return str(s).strip() if s is not None else ""
def pdate(s):
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", s or "")   # 日/月/年
    return f"{int(m.group(3)):04d}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else None

FIELDS = {"BUDGET (RM)": "budget", "TOTAL SPEND": "spend", "COST PER LEAD": "cpl",
          "TOTAL RECEIVED": "received", "ACTUAL CAN USED": "used",
          "TOTAL JUNK": "junk", "TOTAL LINK CLICK": "clicks"}

wb = openpyxl.load_workbook(xlsx, data_only=True)
records = []
for sn in wb.sheetnames:
    line = "Maction" if "Maction" in sn else "Pinyu"
    ws = wb[sn]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    anchors = sorted((r, c) for (r, c), v in cells.items() if norm(v).upper().startswith("BUDGET (RM)"))
    for (r, lc) in anchors:
        vc = lc + 1
        name = next((norm(cells.get((rr, lc))) for rr in range(r - 1, max(0, r - 9), -1)
                     if cells.get((rr, lc)) and "Result" in norm(cells.get((rr, lc)))), None)
        datestr = prev = None
        for rr in range(r - 1, max(0, r - 7), -1):
            v = norm(cells.get((rr, lc)))
            if v.upper().startswith("DATE:") and not datestr: datestr = v
            if "PREVIEW" in v.upper() and not prev: prev = v
        got = {}
        for rr in range(r, r + 13):
            lv = norm(cells.get((rr, lc)))
            for k, short in FIELDS.items():
                if lv.upper().startswith(k): got[short] = cells.get((rr, vc))
        marketer = "Maction"
        pm = re.search(r"PINYU\s*\(([^)]+)\)", (prev or "") + " " + (name or ""), re.I)
        if pm: marketer = pm.group(1).strip()
        elif "Winfred" in (name or ""): marketer = "Winfred"
        elif "Wilson" in (name or ""): marketer = "Wilson"
        elif "Sandy" in (name or ""): marketer = "Sandy"
        vm = re.search(r"\(([AB])\)", name or "")
        records.append({"line": line, "marketer": marketer, "variant": (vm.group(1) if vm else None),
                        "report_date": pdate(datestr), "preview_date": pdate(prev),
                        "budget": got.get("budget"), "spend": got.get("spend"),
                        "received": got.get("received"), "used": got.get("used"),
                        "junk": got.get("junk"), "clicks": got.get("clicks")})

out = build_dashboard_data(records, source=os.path.basename(xlsx))
open(os.path.join(ROOT, "dashboard_data.json"), "w", encoding="utf-8").write(
    json.dumps(out, ensure_ascii=False, indent=1))
print(f"[OK] 解析 {len(records)} 条快照 / {len(out['campaigns'])} 场 -> dashboard_data.json  (报告日 {out['report_date']})")
